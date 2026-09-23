"""Deterministic local protocol templates. No model calls or remote resources."""
from dataclasses import dataclass, field
from datetime import date
from hashlib import sha256
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape

from app.core.config import Settings
from app.schemas.alignment import AlignedSegment
from app.schemas.analysis import MeetingAnalysis


class ExportError(Exception):
    def __init__(self, code, message, status=503):
        super().__init__(message)
        self.code, self.message, self.status = code, message, status


@dataclass
class Block:
    kind: str
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


def protocol_blocks(title: str, analysis: MeetingAnalysis, alignment: list[AlignedSegment],
                    meeting_date: date | None = None) -> list[Block]:
    blocks = [Block("title", "ПРОТОКОЛ СОВЕЩАНИЯ"), Block("body", f"Название: {title}"),
              Block("body", f"Дата: {meeting_date.strftime('%d.%m.%Y') if meeting_date else 'не указана'}"),
              Block("heading", "КРАТКОЕ САММАРИ")]
    summary = analysis.summary
    for heading, points in (("Темы", summary.topics), ("Ключевые обсуждения", summary.key_discussions),
                            ("Проблемы и риски", summary.problems_and_risks),
                            ("КЛЮЧЕВЫЕ РЕШЕНИЯ", summary.decisions)):
        if heading == "КЛЮЧЕВЫЕ РЕШЕНИЯ" and summary.main_action_items:
            blocks.append(Block("subheading", "Основные поручения"))
            for item in summary.main_action_items:
                text = item.description
                if item.condition:
                    text = f"Условное поручение: {text}. Условие: {item.condition}"
                text += f". Ответственный: {item.responsible or 'не указан'}; срок: {item.deadline_raw or 'не указан'}."
                for milestone in item.milestones:
                    text += f" Этап: {milestone.description}; срок: {milestone.deadline_raw or 'не указан'}."
                text += f" [сегменты: {', '.join(map(str, item.source_segment_ids))}]"
                blocks.append(Block("body", text))
        blocks.append(Block("heading" if heading.isupper() else "subheading", heading))
        if not points:
            blocks.append(Block("body", "Не зафиксированы."))
        for point in points:
            condition = getattr(point, "condition", None)
            blocks.append(Block("body", point.text + (f" Условие: {condition}" if condition else "") +
                                f" [сегменты: {', '.join(map(str, point.source_segment_ids))}]"))
    blocks.append(Block("heading", "ПОРУЧЕНИЯ"))
    rows = [["№", "Поручение", "Ответственный", "Срок"]]
    for number, item in enumerate(analysis.action_items, 1):
        details = [item.description]
        if item.condition:
            details.append(f"Условное поручение. Условие: {item.condition}")
        for milestone in item.milestones:
            details.append(f"Этап: {milestone.description}; срок: {milestone.deadline_raw or 'не указан'}")
        details.append(f"Сегменты: {', '.join(map(str, item.source_segment_ids))}")
        details.extend(f"[{e.segment_id}] {e.quote}" for e in item.evidence)
        rows.append([str(number), "\n".join(details), item.responsible or "не указан", item.deadline_raw or "не указан"])
    if analysis.action_items:
        blocks.append(Block("table", rows=rows))
    else:
        blocks.append(Block("body", "Поручения не зафиксированы."))
    blocks.append(Block("heading", "ТРАНСКРИПТ"))
    if not alignment:
        blocks.append(Block("body", "Речь не распознана."))
    for segment in alignment:
        speaker = segment.speaker or ", ".join(segment.speaker_ids) or "говорящий не определён"
        blocks.append(Block("body", f"[{segment.id}] {segment.start:.2f}–{segment.end:.2f} {speaker}: {segment.text}"))
    return blocks


def render_docx(path: Path, blocks: list[Block]) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    for name, size in (("Normal", 10), ("Title", 19), ("Heading 1", 13), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
    styles = {"title": "Title", "heading": "Heading 1", "subheading": "Heading 2", "body": "Normal"}
    for block in blocks:
        if block.kind != "table":
            document.add_paragraph(block.text, styles[block.kind])
            continue
        table = document.add_table(rows=0, cols=4)
        table.style = "Table Grid"
        table.autofit = False
        widths = [Cm(value) for value in (0.9, 8.1, 4, 4)]
        for column, width in zip(table.columns, widths):
            column.width = width
        for row_index, values in enumerate(block.rows):
            row = table.add_row()
            for cell, width, value in zip(row.cells, widths, values):
                cell.width, cell.text = width, value
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = row_index == 0
            if row_index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.save(path)


def render_pdf(path: Path, blocks: list[Block], font: Path) -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

    font_name = "MeetingUnicode" + sha256(str(font.resolve()).encode()).hexdigest()[:12]
    loaded_font = TTFont(font_name, str(font))
    texts = [block.text for block in blocks] + [cell for b in blocks for row in b.rows for cell in row]
    required = {ord(c) for text in texts for c in text if c.isprintable() and not c.isspace()}
    if not required <= loaded_font.face.charToGlyph.keys():
        raise ExportError("pdf_font_glyph_missing", "Выбранный TTF не содержит все символы протокола.")
    pdfmetrics.registerFont(loaded_font)
    body = ParagraphStyle("Body", fontName=font_name, fontSize=10, leading=14, splitLongWords=True, spaceAfter=6)
    styles = {"body": body, "title": ParagraphStyle("Title", parent=body, fontSize=19, leading=24, spaceAfter=14),
              "heading": ParagraphStyle("Heading", parent=body, fontSize=13, leading=18, spaceBefore=12, keepWithNext=True),
              "subheading": ParagraphStyle("Subheading", parent=body, fontSize=11, leading=15, spaceBefore=6, keepWithNext=True)}
    def paragraph(text, style=body):
        return Paragraph(escape(text).replace("\n", "<br/>"), style)
    flow = []
    document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=56, rightMargin=56, topMargin=48, bottomMargin=48)
    for block in blocks:
        if block.kind != "table":
            flow.append(paragraph(block.text, styles[block.kind]))
            continue
        table = LongTable([[paragraph(cell) for cell in row] for row in block.rows],
                         colWidths=[document.width * part for part in (0.06, 0.46, 0.24, 0.24)],
                         repeatRows=1, splitByRow=1, splitInRow=1)
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                  ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("TOPPADDING", (0, 0), (-1, -1), 6),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        flow.extend([table, Spacer(1, 6)])
    document.build(flow)


class LocalExportService:
    def __init__(self, settings: Settings):
        self.settings = settings

    def export(self, meeting_id: str, title: str, analysis: MeetingAnalysis, alignment: list[AlignedSegment],
               meeting_date: date | None = None) -> dict[str, str]:
        directory = self.settings.data_dir / "exports" / meeting_id
        paths = {kind: directory / f"protocol.{kind}" for kind in ("docx", "pdf")}
        temporary = {kind: directory / f".{uuid4().hex}.{kind}" for kind in paths}
        try:
            directory.mkdir(parents=True, exist_ok=True)
            blocks = protocol_blocks(title, analysis, alignment, meeting_date)
            render_docx(temporary["docx"], blocks)
            candidates = [self.settings.pdf_font_path] if self.settings.pdf_font_path else [
                Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
            font = next((p for p in candidates if p and p.is_file()), None)
            if font is None:
                raise ExportError("pdf_font_missing", "Задайте HACKALEM_PDF_FONT_PATH на локальный TTF с поддержкой RU/KZ.")
            render_pdf(temporary["pdf"], blocks, font)
            for kind in paths:
                temporary[kind].replace(paths[kind])
            return {kind: str(path) for kind, path in paths.items()}
        except ImportError:
            raise ExportError("export_dependency_missing", "Установите requirements-export.txt для локальных DOCX/PDF.") from None
        except ExportError:
            raise
        except OSError:
            raise ExportError("export_storage_error", "Не удалось сохранить локальный протокол.", 507) from None
        except Exception:
            raise ExportError("export_failed", "Не удалось сформировать локальный протокол.", 500) from None
        finally:
            for path in temporary.values():
                path.unlink(missing_ok=True)
