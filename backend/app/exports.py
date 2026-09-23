from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from .config import settings


def export_document(title, document, job_id, format):
    path = settings.data_dir / 'exports' / f'{job_id}.{format}'
    path.parent.mkdir(parents=True, exist_ok=True)
    parts = [(title, 'title'), ('Подтверждённый протокол', 'heading')]
    parts += [(s['text'], 'body') for s in document.get('summary', [])]
    parts += [('Поручения', 'heading')]
    for i, action in enumerate(document.get('actions', []), 1):
        due = action.get('due', {}).get('date') or action.get('due_raw') or 'не указан'
        parts.append((f"{i}. {action['title']} — {action.get('assignee_mention') or 'исполнитель не указан'}; срок: {due}. Статус: {action.get('status', 'open')}", 'body'))
    parts += [('Транскрипт', 'heading')]
    for segment in document['segments']:
        time_label = 'Текст' if segment.get('timing_source') == 'manual' else f"{segment['start_ms']//60000:02d}:{(segment['start_ms']//1000)%60:02d}"
        parts.append((f"[{time_label}] {segment.get('speaker', 'SPEAKER')}: {segment['text']}", 'body'))
    temporary = path.with_suffix(path.suffix + '.tmp')
    if format == 'docx':
        doc = Document()
        doc.styles['Normal'].font.name = 'DejaVu Sans'
        for value, kind in parts:
            if kind in ('title', 'heading'):
                doc.add_heading(value, level=0 if kind == 'title' else 1)
            else:
                doc.add_paragraph(value)
        doc.save(str(temporary))
    elif format == 'pdf':
        font = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
        pdfmetrics.registerFont(TTFont('DejaVu', font))
        styles = getSampleStyleSheet()
        body = ParagraphStyle('bodyLocal', fontName='DejaVu', fontSize=10, leading=15, spaceAfter=7)
        heading = ParagraphStyle('headingLocal', parent=body, fontSize=14, leading=20, textColor=colors.HexColor('#17463a'), spaceBefore=12)
        story = [Paragraph(escape(value), heading if kind != 'body' else body) for value, kind in parts]
        SimpleDocTemplate(str(temporary), title=title, author='HackAlem Local', rightMargin=40, leftMargin=40).build(story)
    else:
        raise ValueError('Unsupported export format')
    temporary.replace(path)
    return {'filename': path.name, 'format': format}
