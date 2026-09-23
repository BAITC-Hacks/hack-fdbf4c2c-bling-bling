from io import BytesIO
import json
import os
import wave

from docx import Document
from fastapi.testclient import TestClient
import imageio_ffmpeg
from pypdf import PdfReader
import pytest

from app.core.config import Settings
from app.main import create_app
from app.schemas.diarization import DiarizationSegment
from app.schemas.transcript import TranscriptRaw, TranscriptSegment
from app.services.analysis import LocalMeetingAnalyzer
from app.services.export import ExportError, LocalExportService
from app.services.local_llm import LocalLLMError

TEXT = "Ерлан, дай есеп до пятницы. Выпуск 94%. Қазақша: Әә Ғғ Ққ Ңң Өө Ұұ Үү Һһ Іі."


class STTDouble:
    calls = 0
    def transcribe(self, path):
        self.calls += 1
        assert path.is_file()
        return TranscriptRaw(model="fixture", device="cpu", compute_type="int8", text=TEXT, duration_seconds=1, options={},
            segments=[TranscriptSegment(id=0, start=0, end=1, text=TEXT)])


class DiarizationDouble:
    calls = 0
    def diarize(self, path):
        self.calls += 1
        return [DiarizationSegment(speaker_id="SPEAKER_00", start=0, end=1)]


class LLMDouble:
    """Explicit model double. Upload, database and DOCX/PDF exports are real."""
    calls = 0
    def generate_json(self, **kwargs):
        self.calls += 1
        return json.dumps({"action_items": [dict(description="Подготовить есеп", responsible="Ерлан", deadline_raw="до пятницы",
            source_segment_ids=[0], confidence=0.9, decision_type="assigned", condition=None, milestones=[], evidence=[dict(segment_id=0, quote=TEXT)])],
            "summary": dict(topics=[dict(text="Выпуск 94%", source_segment_ids=[0], evidence=[dict(segment_id=0, quote=TEXT)])],
                key_discussions=[], decisions=[], problems_and_risks=[], main_action_item_indices=[0])}, ensure_ascii=False)


@pytest.fixture
def prepared(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path, database_path=tmp_path / "db.sqlite3", ffmpeg_path=imageio_ffmpeg.get_ffmpeg_exe())
    with TestClient(create_app(settings)) as client:
        meeting_id = client.post("/meetings", json={"title": "Кеңес"}).json()["id"]
        recording = BytesIO()
        with wave.open(recording, "wb") as output:
            output.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
            output.writeframes(b"\0\0" * 16000)
        assert client.post(f"/meetings/{meeting_id}/upload", files={"file": ("test.wav", recording.getvalue(), "audio/wav")}).status_code == 200
        stt, diarizer, llm = STTDouble(), DiarizationDouble(), LLMDouble()
        client.app.state.transcription_service = stt
        client.app.state.diarization_service = diarizer
        client.app.state.meeting_analyzer = LocalMeetingAnalyzer(llm, settings)
        yield client, meeting_id, settings, stt, diarizer, llm


def test_one_inference_to_real_docx_pdf_and_restart(prepared):
    client, meeting_id, settings, stt, diarizer, llm = prepared
    prefix = f"/meetings/{meeting_id}"
    response = client.post(prefix + "/process")
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "ready"
    assert stt.calls == diarizer.calls == llm.calls == 1
    assert client.get(prefix + "/analysis").json()["completed_generations"] == 1
    assert client.get(prefix + "/action-items").json()["prompt_version"] == "joint-1"
    assert client.get(prefix + "/summary").json()["main_action_items"][0]["deadline_raw"] == "до пятницы"
    docx = client.get(prefix + "/exports/docx")
    pdf = client.get(prefix + "/exports/pdf")
    assert docx.status_code == pdf.status_code == 200
    text = "\n".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
    assert TEXT in text and "Ерлан" in text
    document = Document(BytesIO(docx.content))
    assert document.paragraphs[0].text == "ПРОТОКОЛ СОВЕЩАНИЯ"
    assert document.paragraphs[0].style.name == "Title"
    assert "Дата: не указана" in text
    assert [c.text for c in document.tables[0].rows[0].cells] == ["№", "Поручение", "Ответственный", "Срок"]
    assert document.tables[0].cell(1, 2).text == "Ерлан"
    assert document.tables[0].cell(1, 3).text == "до пятницы"
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "94%" in pdf_text and "Қазақша" in pdf_text and "до пятницы" in pdf_text
    for heading in ("ПРОТОКОЛ СОВЕЩАНИЯ", "КРАТКОЕ САММАРИ", "КЛЮЧЕВЫЕ РЕШЕНИЯ", "ПОРУЧЕНИЯ", "ТРАНСКРИПТ"):
        assert heading in text and heading in pdf_text
    assert client.post(prefix + "/process").json() == response.json()
    assert llm.calls == 1
    with TestClient(create_app(settings)) as restarted:
        assert restarted.post(prefix + "/process").json() == response.json()
        assert restarted.get(prefix + "/analysis").status_code == 200


def test_export_retry_reuses_completed_analysis(prepared):
    client, meeting_id, settings, stt, diarizer, llm = prepared
    prefix = f"/meetings/{meeting_id}"
    class UnavailableExport:
        def export(self, *args):
            raise ExportError("pdf_font_missing", "font missing")
    client.app.state.export_service = UnavailableExport()
    assert client.post(prefix + "/process").status_code == 503
    state = client.get(prefix + "/pipeline/status").json()
    assert state["stage"] == "export" and state["status"] == "failed"
    assert client.get(prefix + "/analysis").status_code == 200
    assert client.get(prefix + "/result").status_code == 200
    client.app.state.export_service = LocalExportService(settings)
    assert client.post(prefix + "/process").status_code == 200
    assert stt.calls == diarizer.calls == llm.calls == 1


def test_failed_analysis_is_retryable_without_repeating_audio(prepared):
    client, meeting_id, settings, stt, diarizer, llm = prepared
    prefix = f"/meetings/{meeting_id}"
    class Unavailable:
        def analyze(self, sources):
            raise LocalLLMError("llm_unavailable", "offline")
    client.app.state.meeting_analyzer = Unavailable()
    assert client.post(prefix + "/process").status_code == 503
    assert client.get(prefix + "/action-items").status_code == 409
    assert client.get(prefix + "/pipeline/status").json()["error_code"] == "llm_unavailable"
    client.app.state.meeting_analyzer = LocalMeetingAnalyzer(llm, settings)
    assert client.post(prefix + "/process").status_code == 200
    assert stt.calls == diarizer.calls == llm.calls == 1


def test_concurrent_pipeline_and_legacy_calls_are_blocked(prepared):
    client, meeting_id, settings, _, _, _ = prepared
    prefix = f"/meetings/{meeting_id}"
    class CheckingLLM(LLMDouble):
        def generate_json(self, **kwargs):
            assert client.post(prefix + "/process").status_code == 409
            assert client.post(prefix + "/extract-action-items").status_code == 409
            assert client.post(prefix + "/summarize").status_code == 409
            return super().generate_json(**kwargs)
    llm = CheckingLLM()
    client.app.state.meeting_analyzer = LocalMeetingAnalyzer(llm, settings)
    assert client.post(prefix + "/process").status_code == 200
    assert llm.calls == 1


def test_existing_legacy_results_reused_without_inference(prepared):
    client, meeting_id, settings, _, _, llm = prepared
    prefix = f"/meetings/{meeting_id}"
    # Establish valid compatibility records, then remove only new pipeline cache
    # to emulate a meeting created by the prior two-stage architecture.
    assert client.post(prefix + "/process").status_code == 200
    from app.models.pipeline import MeetingPipeline, MeetingAnalysisRecord
    with client.app.state.session_factory() as session:
        session.delete(session.get(MeetingPipeline, meeting_id))
        session.delete(session.get(MeetingAnalysisRecord, meeting_id))
        session.commit()
    assert client.post(prefix + "/process").status_code == 200
    assert llm.calls == 1
    assert client.get(prefix + "/analysis").json()["mode"] == "reused"


def test_analyze_and_result_never_repeat_inference(prepared):
    client, meeting_id, settings, stt, diarizer, llm = prepared
    prefix = f"/meetings/{meeting_id}"
    assert client.get(prefix + "/result").status_code == 409
    assert stt.calls == diarizer.calls == llm.calls == 0
    assert client.post(prefix + "/analyze").status_code == 200
    result = client.get(prefix + "/result").json()
    assert set(result) == {"meeting", "summary", "action_items", "transcript"}
    assert result["meeting"]["meeting_date"] is None
    assert result["transcript"][0]["speaker_ids"] == ["SPEAKER_00"]
    assert result["transcript"][0]["text"] == TEXT
    assert result["action_items"][0]["responsible"] == "Ерлан"
    for _ in range(3):
        assert client.get(prefix + "/result").json() == result
    assert client.post(prefix + "/analyze").status_code == 200
    assert stt.calls == diarizer.calls == llm.calls == 1
    with TestClient(create_app(settings)) as restarted:
        class Forbidden:
            def __getattr__(self, name):
                raise AssertionError("GET result must not use any AI service")
        for service in ("transcription_service", "diarization_service", "meeting_analyzer", "export_service"):
            setattr(restarted.app.state, service, Forbidden())
        assert restarted.get(prefix + "/result").json() == result
        assert restarted.post(prefix + "/analyze").status_code == 200


def test_result_unknown_and_invalid_ids(prepared):
    from uuid import uuid4
    client = prepared[0]
    assert client.get(f"/meetings/{uuid4()}/result").status_code == 404
    assert client.get("/meetings/bad-id/result").status_code == 422


def test_export_long_table_row_and_explicit_date(prepared):
    from datetime import date
    from pathlib import Path
    from app.schemas.analysis import MeetingAnalysis
    from app.schemas.alignment import AlignedSegment
    client, meeting_id, settings, _, _, llm = prepared
    prefix = f"/meetings/{meeting_id}"
    assert client.post(prefix + "/analyze").status_code == 200
    analysis = MeetingAnalysis.model_validate(client.get(prefix + "/analysis").json())
    # Stress layout with one row longer than a page and literal HTML metacharacters.
    analysis.action_items[0].description = "Начало <script> & есеп " + "длинное поручение " * 400 + "Конец строки"
    alignment = [AlignedSegment.model_validate(s) for s in client.get(prefix + "/result").json()["transcript"]]
    paths = LocalExportService(settings).export("layout-test", "Кеңес <b> & тест", analysis, alignment, date(2026, 9, 20))
    document = Document(paths["docx"])
    assert "Дата: 20.09.2026" in [p.text for p in document.paragraphs]
    assert "Конец строки" in document.tables[0].cell(1, 1).text
    reader = PdfReader(paths["pdf"])
    assert len(reader.pages) > 1
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Конец строки" in text and "20.09.2026" in text and "<script>" in text
    assert llm.calls == 1
    assert all(Path(path).is_file() for path in paths.values())


def test_empty_export_uses_no_invented_content(prepared):
    from app.schemas.analysis import MeetingAnalysis
    from app.schemas.summary import MeetingSummary
    settings = prepared[2]
    empty = MeetingAnalysis(action_items=[], summary=MeetingSummary(topics=[], key_discussions=[], decisions=[],
        problems_and_risks=[], main_action_items=[]), source_segments=[], mode="empty", generation_requests=0,
        completed_generations=0, chunks=0)
    paths = LocalExportService(settings).export("empty", "Без речи", empty, [])
    document = Document(paths["docx"])
    text = "\n".join(p.text for p in document.paragraphs)
    assert not document.tables
    assert "Поручения не зафиксированы." in text and "Речь не распознана." in text


def metric_events(caplog):
    return [record.analysis_metrics for record in caplog.records if hasattr(record, "analysis_metrics")]


def test_analysis_and_cached_reads_log_actual_calls_and_original_model(prepared, caplog):
    caplog.set_level("INFO", logger="app.metrics")
    client, meeting_id, settings, stt, diarizer, llm = prepared
    settings.ollama_model = "fixture-v1"
    prefix = f"/meetings/{meeting_id}"
    assert client.post(prefix + "/analyze").status_code == 200
    settings.ollama_model = "fixture-v2"
    assert client.get(prefix + "/result").status_code == 200
    assert client.post(prefix + "/analyze").status_code == 200
    with TestClient(create_app(settings)) as restarted:
        assert restarted.get(prefix + "/result").status_code == 200
    records = metric_events(caplog)
    assert len(records) == 4
    assert [m["operation"] for m in records] == ["analyze", "result_read", "analyze", "result_read"]
    assert [m["llm_calls"] for m in records] == [1, 0, 0, 0]
    assert [m["cache_status"] for m in records] == ["miss", "hit", "hit", "hit"]
    assert all(m["transcript_characters"] == len(TEXT) and m["model"] == "fixture-v1" for m in records)
    assert records[0]["input_tokens_estimate"] > 0
    assert all(m["input_tokens_estimate"] == 0 for m in records[1:])
    assert all(m["duration_seconds"] >= 0 and m["meeting_id"] == meeting_id for m in records)
    assert stt.calls == diarizer.calls == llm.calls == 1


def test_export_failure_and_retry_metrics_count_current_operation_only(prepared, caplog):
    caplog.set_level("INFO", logger="app.metrics")
    client, meeting_id, settings, _, _, llm = prepared
    class BrokenExport:
        def export(self, *args):
            raise ExportError("pdf_font_missing", "font missing")
    client.app.state.export_service = BrokenExport()
    prefix = f"/meetings/{meeting_id}"
    assert client.post(prefix + "/analyze").status_code == 503
    assert client.get(prefix + "/result").status_code == 200
    client.app.state.export_service = LocalExportService(settings)
    assert client.post(prefix + "/analyze").status_code == 200
    records = metric_events(caplog)
    assert [m["llm_calls"] for m in records] == [1, 0, 0]
    assert [m["cache_status"] for m in records] == ["miss", "hit", "hit"]
    assert [m["status"] for m in records] == ["error", "success", "success"]
    assert llm.calls == 1


def test_unready_result_logs_zero_calls(prepared, caplog):
    caplog.set_level("INFO", logger="app.metrics")
    client, meeting_id, _, _, _, llm = prepared
    assert client.get(f"/meetings/{meeting_id}/result").status_code == 409
    event, = metric_events(caplog)
    assert event["llm_calls"] == event["input_tokens_estimate"] == llm.calls == 0
    assert event["cache_status"] == "miss" and event["status"] == "error"
    assert event["transcript_characters"] is None


def test_metrics_read_legacy_model_without_reprocessing(prepared, caplog):
    from app.models.pipeline import MeetingAnalysisRecord
    client, meeting_id, settings, _, _, llm = prepared
    settings.ollama_model = "original-local-model"
    prefix = f"/meetings/{meeting_id}"
    assert client.post(prefix + "/analyze").status_code == 200
    with client.app.state.session_factory() as session:
        row = session.get(MeetingAnalysisRecord, meeting_id)
        row.result = {key: value for key, value in row.result.items() if key != "model"}
        session.commit()
    caplog.clear()
    caplog.set_level("INFO", logger="app.metrics")
    settings.ollama_model = "different-current-setting"
    assert client.get(prefix + "/result").status_code == 200
    event, = metric_events(caplog)
    assert event["model"] == "original-local-model" and event["cache_status"] == "hit"
    assert event["llm_calls"] == 0 and llm.calls == 1


@pytest.mark.skipif(os.environ.get("HACKALEM_BROWSER_TESTS") != "1", reason="Optional live browser test; see docs/UI.md")
@pytest.mark.parametrize("width", [1280, 390])
def test_browser_upload_process_tabs_reload_and_downloads(prepared, width):
    import re
    import socket
    import threading
    from pathlib import Path
    import uvicorn
    from playwright.sync_api import sync_playwright, expect

    client, _, _, stt, diarizer, llm = prepared
    released = threading.Event()
    class PausedSTT:
        def transcribe(self, path):
            released.wait(10)
            return stt.transcribe(path)
    client.app.state.transcription_service = PausedSTT()
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    base = f"http://127.0.0.1:{sock.getsockname()[1]}"
    server = uvicorn.Server(uvicorn.Config(client.app, lifespan="off", log_level="error"))
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    recording = BytesIO()
    with wave.open(recording, "wb") as output:
        output.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
        output.writeframes(b"\0\0" * 16000)
    qa = Path("data/qa/stage14")
    qa.mkdir(parents=True, exist_ok=True)
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(executable_path=os.environ.get("HACKALEM_TEST_BROWSER"))
            try:
                page = browser.new_page(viewport={"width": width, "height": 1000}, accept_downloads=True)
                errors, requests = [], []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.on("request", lambda request: requests.append((request.method, request.url)))
                page.goto(base)
                expect(page.locator("#analyze")).to_be_disabled()
                page.locator("#title").fill("Кеңес <img src=x onerror=alert(1)>")
                page.locator("#meeting-date").fill("2026-09-20")
                if width == 1280:
                    page.locator("#file").set_input_files({"name": "broken.wav", "mimeType": "audio/wav", "buffer": b"broken"})
                    page.locator("#upload-button").click()
                    expect(page.locator("#error")).to_be_visible()
                    expect(page.locator("#analyze")).to_be_disabled()
                    assert stt.calls == diarizer.calls == llm.calls == 0
                page.locator("#file").set_input_files({"name": "test.wav", "mimeType": "audio/wav", "buffer": recording.getvalue()})
                page.locator("#upload-button").click()
                expect(page.locator("#analyze")).to_be_enabled(timeout=15000)
                page.locator("#analyze").click()
                expect(page.locator("#progress")).to_have_text(re.compile("Распознавание речи"), timeout=5000)
                expect(page.locator("#analyze")).to_be_disabled()
                released.set()
                expect(page.locator("#downloads")).to_be_visible(timeout=15000)
                assert stt.calls == diarizer.calls == llm.calls == 1
                expect(page.locator("#summary")).to_contain_text("Выпуск 94%")
                post_count = sum(method == "POST" for method, _ in requests)
                for _ in range(2):
                    page.locator("#tab-actions").click()
                    expect(page.locator("#actions")).to_contain_text("Ерлан")
                    page.locator("#actions .source").first.click()
                    expect(page.locator("#transcript")).to_be_visible()
                    expect(page.locator("#transcript")).to_contain_text("Қазақша")
                    page.locator("#tab-summary").click()
                page.locator("#tab-actions").click()
                page.screenshot(path=str(qa / f"ui-{width}.png"), full_page=True)
                page.reload()
                expect(page.locator("#downloads")).to_be_visible()
                page.locator("#refresh").click()
                expect(page.locator("#summary")).to_contain_text("Выпуск 94%")
                for kind in ("docx", "pdf"):
                    with page.expect_download() as download:
                        page.locator(f"#{kind}").click()
                    download.value.save_as(str(qa / f"protocol.{kind}"))
                assert sum(method == "POST" for method, _ in requests) == post_count
                assert stt.calls == diarizer.calls == llm.calls == 1
                assert all(url.startswith(base + "/") for _, url in requests)
                assert page.locator("img").count() == 0
                assert not errors
                assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            finally:
                browser.close()
    finally:
        released.set()
        server.should_exit = True
        thread.join(timeout=15)
        sock.close()
    assert not thread.is_alive()
